"""Parser spreadsheet (plan.md §13.1 Spreadsheet Parser).

Setiap sheet menjadi satu halaman, karena ekspor mutasi bank dan laporan POS sering
memuat beberapa sheet dengan makna berbeda.
"""

from __future__ import annotations

import datetime as dt
import io
import zipfile
from decimal import Decimal
from typing import Any

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.parsers.base import DocumentParseError, DocumentParser, ParsedPage, ParseResult

MAX_ROWS_PER_SHEET = 20000


def stringify(value: Any) -> str | None:
    """Mengubah nilai sel menjadi string tanpa merusak presisi angka.

    plan.md §44.14 melarang float untuk uang. openpyxl mengembalikan angka sebagai int
    atau float, jadi nilai float diubah lewat Decimal(repr(...)) agar representasi
    desimalnya tidak melebar menjadi artefak biner seperti 0.1 -> 0.1000000000000000055.
    """
    if value is None:
        return None

    if isinstance(value, bool):
        return "true" if value else "false"

    if isinstance(value, float):
        return format(Decimal(repr(value)).normalize(), "f")

    if isinstance(value, Decimal):
        return format(value.normalize(), "f")

    if isinstance(value, dt.datetime):
        return value.date().isoformat() if value.time() == dt.time.min else value.isoformat()

    if isinstance(value, dt.date):
        return value.isoformat()

    text = str(value).strip()

    return text or None


class SpreadsheetParser(DocumentParser):
    name = "spreadsheet"
    version = "1.0"
    content_kind = "table"
    extensions = ("xlsx", "xlsm", "xls")
    mime_types = (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "application/vnd.ms-excel.sheet.macroenabled.12",
        "application/vnd.ms-excel",
    )

    def parse(self, content: bytes, filename: str) -> ParseResult:
        try:
            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        except (InvalidFileException, zipfile.BadZipFile) as exception:
            # .xlsx adalah kontainer zip. Berkas .xls lama (BIFF) maupun berkas rusak
            # gagal di lapisan zip, dan menolaknya lebih jujur daripada menghasilkan
            # sheet kosong.
            raise DocumentParseError(
                "Format spreadsheet tidak didukung. Simpan ulang sebagai .xlsx."
            ) from exception
        except (OSError, ValueError, KeyError) as exception:
            raise DocumentParseError(f"Spreadsheet tidak dapat dibaca: {exception}") from exception

        pages: list[ParsedPage] = []

        try:
            for index, sheet in enumerate(workbook.worksheets, start=1):
                rows: list[list[str | None]] = []
                truncated = False

                for row_index, row in enumerate(sheet.iter_rows(values_only=True)):
                    if row_index >= MAX_ROWS_PER_SHEET:
                        truncated = True
                        break

                    cells = [stringify(cell) for cell in row]

                    # Baris kosong di akhir sheet adalah artefak umum ekspor Excel.
                    if any(cell is not None for cell in cells):
                        rows.append(cells)

                pages.append(
                    ParsedPage(
                        page_number=index,
                        kind="sheet",
                        label=sheet.title,
                        rows=rows,
                        metadata={
                            "row_count": len(rows),
                            "column_count": max((len(row) for row in rows), default=0),
                            "truncated": truncated,
                        },
                    )
                )
        finally:
            workbook.close()

        if not pages:
            raise DocumentParseError("Spreadsheet tidak memiliki sheet.")

        return self.result(pages, metadata={"sheet_count": len(pages)})
